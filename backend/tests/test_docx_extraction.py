# Copyright (c) 2026 Simon and Handwriting Web contributors.
from pathlib import Path

from docx import Document

from backend.docx_extractor import extract_text_from_docx


def test_extract_text_from_docx_basic(tmp_path):
    doc_path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("第一行测试文本")
    doc.add_paragraph("第二行测试文本\t带制表符")
    doc.save(str(doc_path))

    extracted = extract_text_from_docx(str(doc_path))
    assert "第一行测试文本" in extracted
    assert "第二行测试文本\t带制表符" in extracted


def test_extract_math_symbols_from_generated_docx(tmp_path: Path) -> None:
    """2026-10-15：自行构建数学文本样本，不依赖开发者目录中的私有文档."""
    doc_path = tmp_path / "math.docx"
    formulas = ["LG6={a∣a∈Σ}", "N⇒ND⇒NDD⇒NDDD", "S→1∣3∣5∣7∣9", "E⇒E+T⇒T+T"]
    doc = Document()
    for formula in formulas:
        doc.add_paragraph(formula)
    doc.save(str(doc_path))

    extracted = extract_text_from_docx(str(doc_path))
    for formula in formulas:
        assert formula in extracted
